#!/usr/bin/env python3
"""
Pre-submission checks for the Alzheimer's Research & Therapy package
====================================================================
  1. Figures: each PDF opens, every font is embedded (pdffonts), every raster is at
     least 300 dpi (pdfimages), there is no in-figure title, and Fig. 2 matches a
     fresh render from the bootstrap output. PNG figures are checked for dpi.
  2. Privacy scan of every file: ADNI PTID/RID, OASIS and NACC identifiers, dates,
     identifier-like column names, and tables with one row per participant.
  3. Number cross-check: every Additional-file value that also appears in the
     manuscript must appear there, in context, at the manuscript's precision.
  4. Additional-file titles must match the manuscript's Supplementary information.
  5. With --zip, the package is zipped only if every check passes.

Usage:
    python check_submission_package.py --package <dir> --manuscript <pdf> [--zip <out.zip>]
"""
from pathlib import Path
import argparse
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

import openpyxl
import pandas as pd
from PIL import Image

ROOT = Path(__file__).resolve().parent
POPPLER_FALLBACK = Path("/software/python-anaconda-2022.05-el8-x86_64/envs/R/bin")
OLD_TITLES = ["Cross-cohort decomposition", "prevalence by cohort", "GMM threshold sensitivity",
              "stratified by diagnosis", "Within-OASIS3 continuous"]
COHORT_SIZES = {438, 547, 368, 142, 413, 296}
ID_PATTERNS = {
    "ADNI PTID": r"\b\d{3}_S_\d{4}\b",
    "OASIS ID": r"\bOAS3\d{4}\b|\bOAS\d{5}\b",
    "NACC ID": r"\bNACC\d{6}\b",
}
MONTH = r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*"
DATE_PATTERNS = {
    "ISO date": r"\b\d{4}-\d{2}-\d{2}\b",
    "slash date": r"\b\d{1,2}/\d{1,2}/\d{2,4}\b",
    "compact date": r"\b(?:19|20)\d{2}(?:0[1-9]|1[0-2])(?:0[1-9]|[12]\d|3[01])\b",
    "day-month-year": rf"\b\d{{1,2}}[\s\-]?{MONTH}[\s\-]?\d{{2,4}}\b",
    "month-day-year": rf"\b{MONTH}\.?\s+\d{{1,2}}(?:st|nd|rd|th)?,?\s+\d{{4}}\b",
    "month-year": rf"\b{MONTH}\s+(?:19|20)\d{{2}}\b",
}
ID_HEADER = re.compile(r"(?i)\b(subject|subj|id|ids|ptid|rid|naccid|oasisid|date|dates|visit|visits|session|sessions|examdate|scandate)\b")

failures, notes = [], []


def fail(msg):
    failures.append(msg)
    print("  FAIL:", msg)


def ok(msg):
    print("  ok:  ", msg)


def tool(name):
    found = shutil.which(name)
    return found or str(POPPLER_FALLBACK / name)


def pdf_text(path):
    """Text layer via poppler's pdftotext (keeps ligatures and reading order); Ghostscript as fallback."""
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "t.txt"
        if Path(tool("pdftotext")).exists():
            subprocess.run([tool("pdftotext"), "-layout", str(path), str(out)], check=True, capture_output=True)
        else:
            subprocess.run(["gs", "-q", "-dNOPAUSE", "-dBATCH", "-dSAFER", "-sDEVICE=txtwrite",
                            f"-sOutputFile={out}", str(path)], check=True, capture_output=True)
        return out.read_text(errors="replace")


# ----------------------------------------------------------------------------- 1. figures
def check_figures(pkg):
    print("\n[1] Figures")
    figdir = pkg / "figures"
    for pdf in sorted(figdir.glob("*.pdf")):
        info = subprocess.run([tool("pdfinfo"), str(pdf)], capture_output=True, text=True)
        pages = re.search(r"Pages:\s+(\d+)", info.stdout)
        if info.returncode or not pages:
            fail(f"{pdf.name}: does not open ({info.stderr.strip()})")
            continue
        fonts = subprocess.run([tool("pdffonts"), str(pdf)], capture_output=True, text=True).stdout.splitlines()[2:]
        not_embedded = [l.split()[0] for l in fonts if l.split()[-5] != "yes"]
        types = sorted({" ".join(l.split()[1:3]) for l in fonts})
        if not_embedded:
            fail(f"{pdf.name}: fonts not embedded {not_embedded}")
        imgs = subprocess.run([tool("pdfimages"), "-list", str(pdf)], capture_output=True, text=True).stdout.splitlines()[2:]
        low = [l for l in imgs if min(int(l.split()[12]), int(l.split()[13])) < 300]
        if low:
            fail(f"{pdf.name}: {len(low)} raster image(s) below 300 dpi")
        titles = [t for t in OLD_TITLES if t in pdf_text(pdf)]
        if titles:
            fail(f"{pdf.name}: in-figure title present {titles}")
        if not (not_embedded or low or titles):
            ok(f"{pdf.name}: opens ({pages.group(1)} page), {len(fonts)} font(s) embedded [{', '.join(types)}], "
               f"{len(imgs)} raster image(s), no in-figure title")
    for png in sorted(figdir.glob("*.png")):
        with Image.open(png) as im:
            dpi = im.info.get("dpi")
            msg = f"{png.name}: {im.size[0]}x{im.size[1]} px, dpi metadata {tuple(round(d) for d in dpi) if dpi else 'absent'}"
        if dpi and min(dpi) < 300:
            fail(msg + " (below 300 dpi)")
        else:
            ok(msg + ("" if dpi else "; check the printed width gives at least 300 dpi"))
    fig2 = figdir / "Fig2.pdf"
    if fig2.exists():
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run([sys.executable, str(ROOT / "plot_crosscohort_severity_figure.py"), "--outdir", tmp],
                           check=True, capture_output=True)
            strip = lambda b: re.sub(rb"/CreationDate \(D:[^)]*\)", b"", b)
            fresh = (Path(tmp) / "figure_primary_auc_with_ci.pdf").read_bytes()
        if strip(fresh) == strip(fig2.read_bytes()):
            ok("Fig2.pdf is identical (apart from its timestamp) to a fresh render of the bootstrap output in Additional file 1")
        else:
            fail("Fig2.pdf differs from a fresh render of crosscohort_bootstrap_ci/bootstrap_auc_summary.csv")


# ----------------------------------------------------------------------------- 2. privacy
def scan_text(label, text, allow_dates=False):
    hits = []
    for name, pat in ID_PATTERNS.items():
        hits += [(name, m) for m in re.findall(pat, text)]
    date_hits = [(name, m) for name, pat in DATE_PATTERNS.items() for m in re.findall(pat, text)]
    for name, m in hits:
        fail(f"{label}: {name} {m!r}")
    if date_hits:
        if allow_dates:
            notes.append(f"{label}: dates in manuscript text {sorted({m for _, m in date_hits})}")
        else:
            for name, m in date_hits:
                fail(f"{label}: {name} {m!r}")
    return len(hits) + (0 if allow_dates else len(date_hits))


def check_privacy(pkg, manuscript_names):
    print("\n[2] Privacy scan")
    for path in sorted(p for p in pkg.rglob("*") if p.is_file()):
        rel = path.relative_to(pkg)
        n_before = len(failures)
        if path.suffix == ".xlsx":
            wb = openpyxl.load_workbook(path)
            cells = 0
            for ws in wb.worksheets:
                scan_text(f"{rel}!sheet name", ws.title)
                rows = list(ws.iter_rows(values_only=True))
                if ws.title != "README" and rows:
                    bad = [h for h in rows[0] if h and ID_HEADER.search(re.sub(r"[_\-]", " ", str(h)))]
                    if bad:
                        fail(f"{rel}!{ws.title}: identifier-like column(s) {bad}")
                    n_rows = len(rows) - 1
                    if n_rows > 100 or n_rows in COHORT_SIZES:
                        fail(f"{rel}!{ws.title}: {n_rows} rows (participant-level?)")
                text = "\n".join(str(v) for r in rows for v in r if v is not None)
                cells += sum(v is not None for r in rows for v in r)
                scan_text(f"{rel}!{ws.title}", text)
            with zipfile.ZipFile(path) as z:
                for member in ("docProps/core.xml", "docProps/app.xml"):
                    if member in z.namelist():
                        scan_text(f"{rel}:{member}", z.read(member).decode("utf-8", "replace"))
            if len(failures) == n_before:
                ok(f"{rel}: {len(wb.worksheets)} sheets, {cells} cells and document properties clean")
        elif path.suffix == ".pdf":
            is_ms = path.name in manuscript_names
            scan_text(str(rel), pdf_text(path), allow_dates=is_ms)
            meta = subprocess.run([tool("pdfinfo"), str(path)], capture_output=True, text=True).stdout
            stamps = re.findall(r"(?m)^(CreationDate|ModDate):\s+(.*)$", meta)
            if stamps:
                notes.append(f"{rel}: file-creation timestamp in PDF metadata ({'; '.join(f'{k} {v.strip()}' for k, v in stamps)}), not participant data")
            if len(failures) == n_before:
                ok(f"{rel}: text layer clean" + (" (dates allowed in manuscript text)" if is_ms else ""))
        elif path.suffix in (".md", ".txt", ".csv"):
            scan_text(str(rel), path.read_text(errors="replace"))
            if len(failures) == n_before:
                ok(f"{rel}: clean")
        elif path.suffix in (".png", ".jpg", ".tif", ".tiff"):
            notes.append(f"{rel}: raster image, no text layer to scan; inspect visually")
        else:
            notes.append(f"{rel}: not scanned (unrecognized type)")


# ----------------------------------------------------------------------------- 3-4. manuscript cross-check
def tokens_of(text):
    text = text.replace("−", "-")
    return re.findall(r"(?<![A-Za-z0-9.])-?\d+(?:\.\d+)?", text)


def find_seq(tokens, seq, gap=3):
    best = (0, 0)
    for i, t in enumerate(tokens):
        if t != seq[0]:
            continue
        j, matched = i, 1
        for e in seq[1:]:
            hit = None
            for k in range(j + 1, min(j + 2 + gap, len(tokens))):
                if tokens[k] == e:
                    hit = k
                    break
                if "." in tokens[k]:
                    break
            if hit is None:
                break
            j, matched = hit, matched + 1
        if matched == len(seq):
            return True, None
        best = max(best, (matched, i))
    m, i = best
    return False, " ".join(tokens[max(0, i - 2): i + len(seq) + 4]) if m else "first value not found"


def f3(x):
    return f"{x:.3f}"


def f2(x):
    return f"{x:.2f}"


def pct_token(x):
    v = round(float(x), 1)
    return str(int(v)) if v == int(v) else f"{v:.1f}"


def cross_check(pkg, manuscript):
    print(f"\n[3] Number cross-check against {manuscript.name}")
    text = pdf_text(manuscript)
    toks = tokens_of(text)
    af = pkg / "additional_files"
    read = lambda n, s: pd.read_excel(af / f"Additional_file_{n}.xlsx", sheet_name=s)
    checks = []

    s1 = read(1, "Primary_model_summary").set_index("Cohort")
    for c in ["ADNI tau90", "ADNI tau180", "OASIS3 tau180", "OASIS3 tau90"]:
        r = s1.loc[c]
        tbl = "Table 3" if c != "OASIS3 tau90" else "Table 4"
        checks.append((f"{tbl} row {c}", [str(int(r["N"])), str(int(r["N tau+"]))] + [f3(r[k]) for k in
                       ["AUC reference", "AUC severity-only", "AUC demo-only", "Severity retention", "Demo retention", "CDR drop", "Full drop"]]))
    p = [s1.loc[c] for c in ["ADNI tau90", "ADNI tau180", "OASIS3 tau180"]]
    checks += [
        ("Abstract/3.5 reference AUCs", [f3(r["AUC reference"]) for r in p]),
        ("3.5 severity retention", [f3(r["Severity retention"]) for r in p]),
        ("3.5 demo retention", [f3(r["Demo retention"]) for r in p]),
        ("3.5 full-strip drops", [f3(r["Full drop"]) for r in p]),
        ("3.5 CDR drops", [f3(r["CDR drop"]) for r in p]),
        ("3.8 OASIS3 tau90 AUCs", [f3(s1.loc["OASIS3 tau90", k]) for k in ["AUC reference", "AUC severity-only", "AUC demo-only"]]),
        ("Table 2 post-harmonization row", [f3(s1.loc["ADNI tau180", "AUC demo-only"]), f3(s1.loc["OASIS3 tau180", "AUC demo-only"]),
                                            f3(abs(s1.loc["ADNI tau180", "AUC demo-only"] - s1.loc["OASIS3 tau180", "AUC demo-only"]))]),
    ]
    b = read(1, "Bootstrap_AUC_CI").set_index(["Cohort", "Experiment"])
    for c in ["ADNI tau90", "ADNI tau180", "OASIS3 tau180"]:
        r = b.loc[(c, "Demo-only")]
        checks.append((f"3.2 demo-only CI {c}", [f3(r["AUC"]), "95", f3(r["95% CI lower"]), f3(r["95% CI upper"])]))
    r = b.loc[("OASIS3 tau180", "Severity-only")]
    checks.append(("3.3 severity-only CI OASIS3 tau180", [f3(r["AUC"]), "95", f3(r["95% CI lower"]), f3(r["95% CI upper"])]))
    d = read(1, "Bootstrap_contrasts").set_index(["Cohort", "Contrast"])
    for c in ["ADNI tau90", "ADNI tau180", "OASIS3 tau180"]:
        for con in ["Reference minus demo-only", "Reference minus severity-only"]:
            r = d.loc[(c, con)]
            checks.append((f"3.5 {con} CI {c}", [f3(r["AUC difference"]), "95", f3(r["95% CI lower"]), f3(r["95% CI upper"])]))

    m2 = read(2, "Recovered_mappings")
    m2 = m2[m2["Cohort"] == "OASIS3 tau180"].set_index("Harmonized variable")
    seq = []
    for v in ["age", "APOE e4 count", "education", "sex", "CDR-SB", "FAQ total"]:
        seq += [str(int(m2.loc[v, "N non-missing"])), str(int(m2.loc[v, "N"]))]
    checks.append(("2.3 OASIS3 tau180 non-missing counts", seq))

    n4 = read(4, "NACC_summary").set_index("Cohort").loc["NACC main"]
    checks.append(("Table 4 row NACC strict A/T", [str(int(n4["N"])), str(int(n4["N A+/T+"]))] + [f3(n4[k]) for k in
                   ["AUC reference", "AUC severity-only", "AUC demo-only", "Severity retention", "Demo retention", "CDR drop", "Full drop"]]))
    nb = read(4, "NACC_feature_blocks")
    nb = nb[nb["Cohort"] == "NACC main"].set_index("Experiment")["AUC"]
    tr = read(4, "Tau_tracer_prevalence")
    checks += [
        ("3.8 NACC AUCs", [f3(nb["Reference-combined"]), f3(nb["Severity-only"]), f3(nb["Demo-only"])]),
        ("3.8 NACC tracer prevalence", [f"{tr['% A+/T+'].iloc[0]:.1f}", f"{tr['% A+/T+'].iloc[1]:.1f}"]),
        ("3.8 NACC demo-only variants", [f3(nb["Demo-only without tau tracer"]), f3(nb["Demo-only, five harmonized ADNI/OASIS3 variables"])]),
        ("3.8 NACC drops", [f3(n4["CDR drop"]), f3(n4["Drop minus CDR, FAQ and MoCA"])]),
        ("2.2 NACC tracer counts", [str(int(tr["N"].iloc[0])), str(int(tr["N"].iloc[1]))]),
    ]

    s5 = read(5, "Table_S5")
    for _, r in s5.iterrows():
        seq = [pct_token(r["% from GMM midpoint"]), f3(r["Threshold (SUVR)"]), str(int(r["N"])), str(int(r["N tau+"]))]
        if pd.notna(r["AUC demo-only"]):
            seq.append(f3(r["AUC demo-only"]))
        checks.append((f"Table S5 row {r['Cohort']} {pct_token(r['% from GMM midpoint'])}%", seq))
    sw = read(5, "Threshold_sweep")
    rel = sw[(sw["Cohort"] == "OASIS3 tau180") & (sw["Reliable"] == "yes") & (sw["Threshold type"] != "published OASIS-3 cutoff")]
    gmm = sw[sw["Threshold type"] == "GMM midpoint"].set_index("Cohort")
    checks += [
        ("3.4 tau180 sweep AUC range", [f3(rel["AUC demo-only"].min()), f3(rel["AUC demo-only"].max())]),
        ("3.4 tau180 GMM midpoint", [f3(gmm.loc["OASIS3 tau180", "Threshold (SUVR)"]), f3(gmm.loc["OASIS3 tau180", "AUC demo-only"]),
                                     str(int(gmm.loc["OASIS3 tau180", "N"])), str(int(gmm.loc["OASIS3 tau180", "N tau+"]))]),
        ("3.4 tau90 GMM midpoint", [f3(gmm.loc["OASIS3 tau90", "Threshold (SUVR)"]), str(int(gmm.loc["OASIS3 tau90", "N tau+"])),
                                    f3(gmm.loc["OASIS3 tau90", "AUC demo-only"])]),
    ]
    pc = read(5, "Published_cutoff_comparison").set_index(["Cohort", "Definition"])
    p180, p90 = pc.loc[("OASIS3 tau180", "Published OASIS-3 cutoff")], pc.loc[("OASIS3 tau90", "Published OASIS-3 cutoff")]
    g180 = pc.loc[("OASIS3 tau180", "GMM midpoint")]
    checks += [
        ("2.2 published cutoff prevalence", [f"{p180['% tau+']:.1f}", str(int(p180["N tau+"])), str(int(p180["N"])),
                                             f"{p90['% tau+']:.1f}", str(int(p90["N tau+"])), str(int(p90["N"]))]),
        ("2.2 amyloid-positive percentage", [f"{p180['% amyloid+ (all with amyloid PET)']:.1f}"]),
        ("2.2 amyloid-negative among tau+", [f"{p180['% amyloid- among tau+']:.1f}", str(int(p180["N tau+ amyloid-"])),
                                             str(int(p180["N tau+ with amyloid PET"])), f"{g180['% amyloid- among tau+']:.1f}",
                                             str(int(g180["N tau+ amyloid-"])), str(int(g180["N tau+ with amyloid PET"]))]),
        ("3.4 published cutoff demo AUCs", [f3(p180["AUC demo-only"]), str(int(p180["N tau+"])), f3(p90["AUC demo-only"]), str(int(p90["N tau+"]))]),
        ("3.4 CDR-SB = 0 among tau+ and severity AUC", [f"{p180['% CDR-SB = 0 among tau+']:.1f}", f"{g180['% CDR-SB = 0 among tau+']:.1f}",
                                                        f3(p180["AUC severity-only"])]),
    ]

    s6 = read(6, "Table_S6").set_index("Term")
    for term, label in [("CDR-SB", "CDR-SB"), ("Age", "Age"), ("APOE e4 carrier", "APOE carrier"), ("APOE e4 count", "APOE count"),
                        ("Education", "Education"), ("Sex: female vs. male (difference of one-hot terms)", "Sex contrast"), ("FAQ", "FAQ")]:
        checks.append((f"Table S6 row {label}", [f2(s6.loc[term, c]) for c in ["ADNI tau90", "ADNI tau180", "OASIS3 tau180"]]))
    checks += [
        ("4.2 CDR-SB coefficient range", [f2(s6.loc["CDR-SB"].min()), f2(s6.loc["CDR-SB"].max())]),
        ("4.2 ADNI APOE count range", [f2(s6.loc["APOE e4 count", "ADNI tau90"]), f2(s6.loc["APOE e4 count", "ADNI tau180"])]),
        ("4.2 OASIS3 age, APOE count, CDR-SB", [f2(s6.loc["Age", "OASIS3 tau180"]), f2(s6.loc["APOE e4 count", "OASIS3 tau180"]),
                                                f2(s6.loc["CDR-SB", "OASIS3 tau180"])]),
    ]

    dx = read(7, "Stratified_AUC").set_index(["Cohort", "Diagnosis stratum"])
    for c in ["ADNI tau90", "ADNI tau180", "OASIS3 tau180"]:
        checks.append((f"3.7 diagnosis counts {c}", [str(int(dx.loc[(c, s), "N"])) for s in ["CN", "MCI", "Dementia"]]))
    o_cn, o_de = dx.loc[("OASIS3 tau180", "CN")], dx.loc[("OASIS3 tau180", "Dementia")]
    a90, a180 = dx.loc[("ADNI tau90", "Dementia")], dx.loc[("ADNI tau180", "Dementia")]
    m90, m180 = dx.loc[("ADNI tau90", "MCI")], dx.loc[("ADNI tau180", "MCI")]
    checks += [
        ("3.7 not-estimated strata N+", [str(int(dx.loc[("ADNI tau90", "CN"), "N tau+"])), str(int(dx.loc[("OASIS3 tau180", "MCI"), "N tau+"]))]),
        ("3.7 OASIS3 CN", [str(int(o_cn["N tau+"])), f3(o_cn["AUC demo-only"]), f3(o_cn["AUC reference"])]),
        ("3.7 OASIS3 Dementia", [str(int(o_de["N tau+"])), str(int(o_de["N tau-"])), f3(o_de["AUC reference"]), f3(o_de["AUC demo-only"])]),
        ("3.7 ADNI Dementia class sizes", [str(int(a90["N tau+"])), str(int(a180["N tau+"])), str(int(a90["N tau-"])), str(int(a180["N tau-"]))]),
        ("3.7 ADNI tau90 Dementia", [f3(a90["AUC reference"]), f3(a90["AUC demo-only"])]),
        ("3.7 ADNI tau180 Dementia", [f3(a180["AUC reference"]), f3(a180["AUC demo-only"])]),
        ("3.7 ADNI MCI", [str(int(m90["N tau+"])), str(int(m180["N tau+"])), f3(m180["AUC demo-only"]), f3(m90["AUC demo-only"]),
                          f3(m180["AUC reference"]), f3(m90["AUC reference"])]),
    ]

    s7 = read(8, "Table_S7").set_index("Cohort")
    for c in ["OASIS3 tau180", "OASIS3 tau90"]:
        checks.append((f"Table S7 row {c}", [f3(s7.loc[c, k]) for k in ["R2 reference", "R2 demo-only", "R2 severity-only", "R2 minus all CDR", "R2 minus FAQ"]]))
    checks.append(("3.4 continuous minus-CDR-SB", [f3(s7.loc["OASIS3 tau180", "R2 reference"]), f3(s7.loc["OASIS3 tau180", "R2 minus all CDR"]),
                                                   f3(s7.loc["OASIS3 tau90", "R2 reference"]), f3(s7.loc["OASIS3 tau90", "R2 minus all CDR"])]))

    n_values = 0
    for label, seq in checks:
        found, near = find_seq(toks, seq)
        n_values += len(seq)
        if found:
            ok(f"{label}: {' '.join(seq)}")
        else:
            fail(f"{label}: Additional-file values {' '.join(seq)} not found in manuscript; closest text: {near}")
    print(f"  {len(checks)} checks covering {n_values} values")

    print(f"\n[4] Additional-file titles against {manuscript.name}")
    lines = []
    for line in text.splitlines():
        s = line.strip()
        if re.fullmatch(r"\d{3,4}", s):
            continue
        lines.append(re.sub(r"\s+\d{3,4}$", "", re.sub(r"^\d{3,4}\s+", "", s)))
    flat = re.sub(r"[^a-z0-9]", "", " ".join(lines).lower())
    for n in range(1, 9):
        title = openpyxl.load_workbook(af / f"Additional_file_{n}.xlsx")["README"]["A2"].value
        if re.sub(r"[^a-z0-9]", "", f"additional file {n}: {title}".lower()) in flat:
            ok(f"Additional file {n}: title matches")
        else:
            fail(f"Additional file {n}: title {title!r} not found in the Supplementary information paragraph")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--package", required=True)
    ap.add_argument("--manuscript", required=True)
    ap.add_argument("--zip")
    args = ap.parse_args()
    pkg, ms = Path(args.package), Path(args.manuscript)
    check_figures(pkg)
    check_privacy(pkg, {ms.name})
    cross_check(pkg, ms)
    print("\nNotes:")
    for n in notes:
        print("  -", n)
    print(f"\nRESULT: {'PASS' if not failures else f'FAIL ({len(failures)} problem(s))'}")
    if args.zip and not failures:
        out = Path(args.zip)
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            for p in sorted(pkg.rglob("*")):
                if p.is_file():
                    z.write(p, Path(pkg.name) / p.relative_to(pkg))
        print(f"zipped: {out}")
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
